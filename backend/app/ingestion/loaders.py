"""Document loaders: extract text (and light structure) from many file types.

Each loader returns a list of ``LoadedSection`` objects. A "section" is a
natural structural unit for the format — a PDF page, a docx paragraph run, a
CSV row-batch, an epub chapter, a markdown/html document. Keeping sections
separate lets the chunker attach meaningful location labels (page numbers,
chapter titles) so large textbooks and bulk data stay navigable after indexing.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable


@dataclass
class LoadedSection:
    """A structural unit of a document with location metadata."""

    text: str
    # Location label shown to users and the LLM, e.g. "p. 12" or "Chapter 3".
    location: str
    # Extra structured metadata merged into every chunk from this section.
    meta: dict = field(default_factory=dict)


class UnsupportedFileType(Exception):
    pass


# --------------------------------------------------------------------------
# Individual format loaders
# --------------------------------------------------------------------------
def load_txt(path: Path) -> list[LoadedSection]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [LoadedSection(text=text, location="full", meta={})]


def load_markdown(path: Path) -> list[LoadedSection]:
    # Treat markdown as text; headings are preserved so the chunker and LLM can
    # use them as natural context.
    text = path.read_text(encoding="utf-8", errors="replace")
    return [LoadedSection(text=text, location="full", meta={"format_detail": "markdown"})]


def load_pdf(path: Path) -> list[LoadedSection]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections: list[LoadedSection] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            sections.append(
                LoadedSection(text=text, location=f"p. {i}", meta={"page": i})
            )
    return sections


def load_docx(path: Path) -> list[LoadedSection]:
    from docx import Document

    doc = Document(str(path))
    # Group paragraphs under their most recent heading so sections are coherent.
    sections: list[LoadedSection] = []
    current_heading = "Introduction"
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if buffer:
            sections.append(
                LoadedSection(
                    text="\n".join(buffer),
                    location=current_heading,
                    meta={"heading": current_heading},
                )
            )
            buffer = []

    for para in doc.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading") and para.text.strip():
            flush()
            current_heading = para.text.strip()
        elif para.text.strip():
            buffer.append(para.text)
    flush()

    # Also capture tables as pipe-delimited text.
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            sections.append(
                LoadedSection(
                    text="\n".join(rows),
                    location=f"Table {t_idx}",
                    meta={"table": t_idx},
                )
            )
    return sections


def load_csv(path: Path, rows_per_section: int = 100) -> list[LoadedSection]:
    """Load a CSV as batches of rows so huge datasets chunk predictably.

    The header is repeated at the top of every batch so each chunk is
    self-describing when retrieved out of context.
    """
    sections: list[LoadedSection] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        try:
            header = next(reader)
        except StopIteration:
            return sections
        header_line = " | ".join(header)

        batch: list[str] = []
        start_row = 1
        for row_num, row in enumerate(reader, start=1):
            batch.append(" | ".join(row))
            if len(batch) >= rows_per_section:
                sections.append(
                    _csv_section(header_line, batch, start_row, row_num)
                )
                batch = []
                start_row = row_num + 1
        if batch:
            sections.append(
                _csv_section(header_line, batch, start_row, start_row + len(batch) - 1)
            )
    return sections


def _csv_section(header: str, rows: list[str], start: int, end: int) -> LoadedSection:
    body = header + "\n" + "\n".join(rows)
    return LoadedSection(
        text=body,
        location=f"rows {start}-{end}",
        meta={"row_start": start, "row_end": end},
    )


def load_epub(path: Path) -> list[LoadedSection]:
    import ebooklib
    from bs4 import BeautifulSoup
    from ebooklib import epub

    book = epub.read_epub(str(path))
    sections: list[LoadedSection] = []
    for i, item in enumerate(book.get_items_of_type(ebooklib.ITEM_DOCUMENT), start=1):
        soup = BeautifulSoup(item.get_content(), "lxml")
        # Prefer a chapter title from the first heading if present.
        heading_tag = soup.find(["h1", "h2", "h3"])
        title = heading_tag.get_text(strip=True) if heading_tag else f"Section {i}"
        text = soup.get_text(separator="\n", strip=True)
        if text.strip():
            sections.append(
                LoadedSection(text=text, location=title, meta={"chapter": title})
            )
    return sections


def load_html(path: Path) -> list[LoadedSection]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_bytes(), "lxml")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    title_tag = soup.find("title")
    title = title_tag.get_text(strip=True) if title_tag else "document"
    text = soup.get_text(separator="\n", strip=True)
    return [LoadedSection(text=text, location=title, meta={"title": title})]


# --------------------------------------------------------------------------
# Dispatch
# --------------------------------------------------------------------------
_LOADERS: dict[str, Callable[[Path], list[LoadedSection]]] = {
    ".txt": load_txt,
    ".text": load_txt,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
    ".epub": load_epub,
    ".html": load_html,
    ".htm": load_html,
}

SUPPORTED_EXTENSIONS = tuple(sorted(_LOADERS.keys()))


def load_document(path: Path) -> list[LoadedSection]:
    """Load any supported document into a list of sections.

    Raises UnsupportedFileType for extensions we don't handle (e.g. legacy
    binary ``.doc`` — see note below).
    """
    ext = path.suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    return loader(path)
